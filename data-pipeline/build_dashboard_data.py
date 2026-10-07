import json
import os
import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path(os.environ.get("CRM_SOURCE_XLSX", "UIO客戶數據平台.xlsx"))
OUTPUT_DIR = Path(os.environ.get("CRM_OUTPUT_DIR", Path(__file__).resolve().parents[1] / "public" / "data"))
AS_OF = date(2026, 10, 2)


def clean_text(value):
    if value is None:
        return ""
    return str(value).strip()


def clean_plate(value):
    return re.sub(r"[^A-Z0-9]", "", clean_text(value).upper())


def clean_phone(*values):
    for value in values:
        digits = re.sub(r"\D", "", clean_text(value))
        if digits:
            return digits
    return ""


def as_number(value):
    if value is None or value == "":
        return None
    try:
        return float(str(value).replace(",", ""))
    except (TypeError, ValueError):
        return None


def as_int(value):
    number = as_number(value)
    return int(round(number)) if number is not None else None


def as_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip().replace("/", "-")
        for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                pass
    return None


def iso(value):
    parsed = as_date(value)
    return parsed.isoformat() if parsed else ""


def add_years(value, years):
    try:
        return value.replace(year=value.year + years)
    except ValueError:
        return value.replace(month=2, day=28, year=value.year + years)


def warranty_state(registration, explicit_expiry, mileage):
    reg = as_date(registration)
    expiry = as_date(explicit_expiry)
    if expiry and expiry.year < 1990:
        expiry = None
    if not expiry and reg:
        expiry = add_years(reg, 3)

    days_left = (expiry - AS_OF).days if expiry else None
    km_left = 100000 - mileage if mileage is not None else None

    expired_time = days_left is not None and days_left <= 0
    expired_mileage = mileage is not None and mileage >= 100000
    if expired_time and expired_mileage:
        status = "已過保（年限與里程）"
    elif expired_time:
        status = "已過保（年限）"
    elif expired_mileage:
        status = "已過保（里程）"
    elif days_left is None or mileage is None:
        status = "待確認"
    elif days_left <= 180 and mileage >= 80000:
        status = "即將過保（雙重）"
    elif days_left <= 180:
        status = "即將過保（年限）"
    elif mileage >= 80000:
        status = "即將過保（里程）"
    else:
        status = "保固內"
    return status, days_left, km_left, expiry


def recommended_action(status, days_since_visit, mileage, tires, battery):
    if status.startswith("即將過保"):
        return "3 日內聯絡並安排保固健檢"
    if status.startswith("已過保") and days_since_visit is not None and days_since_visit >= 365:
        return "啟動久未回廠關懷與車況健檢"
    if mileage is not None and mileage >= 90000:
        return "安排 100,000 km 前預防健檢"
    if days_since_visit is not None and days_since_visit >= 270:
        return "本週由原服務專員邀約回廠"
    if not tires:
        return "確認輪胎年限與磨耗"
    if not battery:
        return "確認電瓶健康度"
    return "依保養週期持續追蹤"


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)

    total_by_plate = {}
    total_sheet = workbook["總表"]
    for row in total_sheet.iter_rows(min_row=2, values_only=True):
        plate = clean_plate(row[0])
        if plate:
            total_by_plate[plate] = row

    records = []
    models = Counter()
    counties = Counter()
    branches = Counter()
    advisors = Counter()
    technicians = Counter()
    sales_representatives = Counter()
    warranty_counts = Counter()
    uio_sheet = workbook["UIO"]

    for row_no, row in enumerate(uio_sheet.iter_rows(min_row=2, values_only=True), 2):
        plate = clean_plate(row[2])
        if not plate:
            continue
        total = total_by_plate.get(plate)

        latest_date_values = [as_date(row[44]), as_date(row[48]), as_date(row[53]), as_date(row[58])]
        latest_dates = [d for d in latest_date_values if d]
        last_visit = max(latest_dates) if latest_dates else None
        mileage_values = [as_int(row[45]), as_int(row[49]), as_int(row[54]), as_int(row[59])]
        mileage_values = [m for m in mileage_values if m is not None and 0 <= m <= 2_000_000]
        mileage = max(mileage_values) if mileage_values else None

        warranty_status, days_left, km_left, warranty_expiry = warranty_state(row[11], row[12], mileage)
        days_since_visit = (AS_OF - last_visit).days if last_visit else None
        owner = clean_text(row[24]) or clean_text(row[34])
        phone = clean_phone(row[30], row[28], row[29], row[40], row[38], row[39])
        email = clean_text(row[31]) or clean_text(row[41])
        address = clean_text(row[33]) or clean_text(row[43])

        vehicle_age = as_number(total[3]) if total else None
        customer_age = as_int(total[2]) if total else None
        county = clean_text(total[5]) if total else ""
        district = clean_text(total[6]) if total else ""
        spend_3m = as_number(total[15]) if total else None
        return_count = as_int(total[12]) if total else None
        tires = as_int(total[29]) if total else None
        battery = as_int(total[30]) if total else None
        advisor = (clean_text(total[10]) if total else "") or clean_text(row[50])
        technician = clean_text(total[11]) if total else ""
        sales_representative = clean_text(row[21])
        branch = clean_text(row[47]) or clean_text(row[52]) or clean_text(row[57])
        model = clean_text(row[10])

        record = {
            "id": row_no - 1,
            "plate": plate,
            "vin": clean_text(row[8]),
            "owner": owner,
            "gender": clean_text(row[25]) or (clean_text(total[1]) if total else ""),
            "customerAge": customer_age,
            "phone": phone,
            "email": email,
            "postalCode": clean_text(row[32]),
            "address": address,
            "county": county,
            "district": district,
            "modelCode": clean_text(row[9]),
            "model": model,
            "registrationDate": iso(row[11]),
            "vehicleAge": round(vehicle_age, 1) if vehicle_age is not None else None,
            "mileage": mileage,
            "lastVisit": last_visit.isoformat() if last_visit else "",
            "daysSinceVisit": days_since_visit,
            "branch": branch,
            "advisor": advisor,
            "salesRepresentative": sales_representative,
            "technician": technician,
            "memberPassport": clean_text(total[9]) if total else "",
            "ownerAuthorization": clean_text(total[8]) if total else "",
            "spend3m": round(spend_3m, 0) if spend_3m is not None else None,
            "returnCount1y": return_count,
            "maintenanceCount1y": as_int(total[12]) if total else None,
            "returnCategory": clean_text(total[13]) if total else "",
            "maintenanceFrequency": clean_text(total[14]) if total else "",
            "beautySpend1y": as_int(total[16]) if total else None,
            "beautyHighSpend": clean_text(total[17]) if total else "",
            "fakeBeautySpend1y": as_int(total[18]) if total else None,
            "windshieldCoatingCount1y": as_int(total[19]) if total else None,
            "windshieldCoatingSpend1y": as_int(total[20]) if total else None,
            "waxCount1y": as_int(total[21]) if total else None,
            "waxSpend1y": as_int(total[22]) if total else None,
            "bodyBeautyCount1y": as_int(total[23]) if total else None,
            "bodyBeautySpend1y": as_int(total[24]) if total else None,
            "interiorCleaningCount1y": as_int(total[25]) if total else None,
            "interiorCleaningSpend1y": as_int(total[26]) if total else None,
            "bodyCoatingCount1y": as_int(total[27]) if total else None,
            "bodyCoatingSpend1y": as_int(total[28]) if total else None,
            "warrantyStatus": warranty_status,
            "warrantyExpiry": warranty_expiry.isoformat() if warranty_expiry else "",
            "warrantyDaysLeft": days_left,
            "warrantyKmLeft": km_left,
            "extendedWarrantyExpiry": iso(row[62]),
            "tireChanges1y": tires,
            "batteryChanges1y": battery,
            "airFilterChanges1y": as_int(total[31]) if total else None,
            "cabinFilterChanges1y": as_int(total[32]) if total else None,
            "wiperChanges1y": as_int(total[33]) if total else None,
            "brakePadChanges1y": as_int(total[34]) if total else None,
            "beltChanges1y": as_int(total[35]) if total else None,
            "sparkPlugChanges1y": as_int(total[36]) if total else None,
            "tireLastChanged": iso(total[37]) if total else "",
            "batteryLastChanged": iso(total[38]) if total else "",
            "airFilterLastChanged": iso(total[39]) if total else "",
            "cabinFilterLastChanged": iso(total[40]) if total else "",
            "wiperLastChanged": iso(total[41]) if total else "",
            "brakePadLastChanged": iso(total[42]) if total else "",
            "beltLastChanged": iso(total[43]) if total else "",
            "sparkPlugLastChanged": iso(total[44]) if total else "",
            "contactAuthorized": clean_text(row[26]),
            "doNotContact": bool(clean_text(row[0])),
            "nextBestAction": recommended_action(warranty_status, days_since_visit, mileage, tires, battery),
        }
        records.append(record)
        if model:
            models[model] += 1
        if county:
            counties[county] += 1
        if branch:
            branches[branch] += 1
        if advisor:
            advisors[advisor] += 1
        if technician:
            technicians[technician] += 1
        if sales_representative:
            sales_representatives[sales_representative] += 1
        warranty_counts[warranty_status] += 1

    workbook.close()

    with (OUTPUT_DIR / "vehicles.json").open("w", encoding="utf-8") as handle:
        json.dump(records, handle, ensure_ascii=False, separators=(",", ":"))

    metadata = {
        "source": os.environ.get("CRM_SOURCE_LABEL", SOURCE.name),
        "asOf": AS_OF.isoformat(),
        "recordCount": len(records),
        "models": [name for name, _ in models.most_common()],
        "counties": [name for name, _ in counties.most_common()],
        "branches": [name for name, _ in branches.most_common()],
        "advisors": [name for name, _ in advisors.most_common()],
        "technicians": [name for name, _ in technicians.most_common()],
        "salesRepresentatives": [name for name, _ in sales_representatives.most_common()],
        "warrantyCounts": dict(warranty_counts),
    }
    with (OUTPUT_DIR / "metadata.json").open("w", encoding="utf-8") as handle:
        json.dump(metadata, handle, ensure_ascii=False, indent=2)

    print(json.dumps({"records": len(records), "bytes": (OUTPUT_DIR / "vehicles.json").stat().st_size, "warranty": dict(warranty_counts)}, ensure_ascii=True))


if __name__ == "__main__":
    main()
