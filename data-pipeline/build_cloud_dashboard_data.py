"""Build dashboard data from the de-identified (雲端_去識別化) UIO workbook.

Personal fields (plate, VIN, owner/driver names, phones, email, address, postal code,
birthdays, village) are never written to the output. Each vehicle keeps only its
識別碼 (e.g. ES-00001) so exported lists can be matched back to the on-prem file.
"""
import json
import os
import re
import sys
from collections import Counter
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent))
import build_dashboard_data as base  # noqa: E402  (shared parsing + warranty/action rules)

SOURCE = Path(os.environ.get("CRM_SOURCE_XLSX", "UIO客戶數據平台_雲端_去識別化.xlsx"))
SOURCE_LABEL = os.environ.get("CRM_SOURCE_LABEL", "UIO客戶數據平台(20261007)_雲端_去識別化.xlsx")
OUTPUT_DIR = base.OUTPUT_DIR
AS_OF = date(2026, 10, 7)
base.AS_OF = AS_OF

CODE_PATTERN = re.compile(r"^ES-\d+$")
# UIO sheet columns that hold personal data in the on-prem file; in this file they must be blank or a code.
UIO_PII_COLUMNS = {2: "牌照號碼", 6: "車身#", 7: "引擎#", 8: "VIN#", 24: "車主", 28: "車主電話", 29: "車主電話2", 30: "車主行動電話", 31: "車主Email", 33: "車主地址", 34: "駕駛", 38: "駕駛電話", 39: "駕駛電話2", 40: "駕駛行動電話", 41: "駕駛Email", 43: "駕駛地址"}


def vehicle_age_years(value):
    """總表車齡 is text like '5年3個月'; fall back to a plain number."""
    text = base.clean_text(value)
    match = re.match(r"^(\d+)年(?:(\d+)個月)?$", text)
    if match:
        return round(int(match.group(1)) + int(match.group(2) or 0) / 12, 1)
    number = base.as_number(value)
    return round(number, 1) if number is not None else None


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)

    total_by_code = {}
    for row in workbook["總表"].iter_rows(min_row=2, values_only=True):
        code = base.clean_text(row[0])
        if code:
            total_by_code[code] = row

    records, seen_codes = [], set()
    models, counties, branches, advisors, technicians, sales_representatives, warranty_counts = (Counter() for _ in range(7))
    leaks = Counter()
    duplicates = unmatched = 0

    for row_no, row in enumerate(workbook["UIO"].iter_rows(min_row=2, values_only=True), 2):
        code = base.clean_text(row[2])
        if not code:
            continue
        for index, label in UIO_PII_COLUMNS.items():
            value = base.clean_text(row[index])
            if value and not CODE_PATTERN.match(value):
                leaks[label] += 1
        if code in seen_codes:
            duplicates += 1
        seen_codes.add(code)
        total = total_by_code.get(code)
        if total is None:
            unmatched += 1
        t = (lambda i: total[i]) if total else (lambda i: None)

        latest_dates = [d for d in (base.as_date(row[44]), base.as_date(row[48]), base.as_date(row[53]), base.as_date(row[58])) if d]
        last_visit = max(latest_dates) if latest_dates else None
        mileage_values = [m for m in (base.as_int(row[45]), base.as_int(row[49]), base.as_int(row[54]), base.as_int(row[59])) if m is not None and 0 <= m <= 2_000_000]
        mileage = max(mileage_values) if mileage_values else None
        warranty_status, days_left, km_left, warranty_expiry = base.warranty_state(row[11], row[12], mileage)
        days_since_visit = (AS_OF - last_visit).days if last_visit else None

        tires, battery = base.as_int(t(34)), base.as_int(t(35))
        advisor = base.clean_text(t(15)) or base.clean_text(row[50])
        technician = base.clean_text(t(16))
        sales_representative = base.clean_text(row[21])
        branch = base.clean_text(row[47]) or base.clean_text(row[52]) or base.clean_text(row[57])
        # 車型 uses the short series code from 總表 (P15, T32…); the full code (P15FVA…) stays in modelCode.
        model_code = base.clean_text(row[9]) or base.clean_text(row[10])
        model = base.clean_text(t(8)) or model_code[:3]
        county = base.clean_text(t(10)) or base.clean_text(row[17])
        district = base.clean_text(t(11))
        spend_3m = base.as_number(t(20))

        records.append({
            "id": row_no - 1,
            "code": code,
            "hasMobile": bool(base.clean_text(row[30]) or base.clean_text(row[40])),
            "hasEmail": bool(base.clean_text(row[31]) or base.clean_text(row[41])),
            "gender": base.clean_text(row[25]) or base.clean_text(t(5)),
            "customerAge": base.as_int(t(6)),
            "county": county,
            "district": district,
            "modelCode": model_code,
            "model": model,
            "registrationDate": base.iso(row[11]),
            "vehicleAge": vehicle_age_years(t(7)),
            "mileage": mileage,
            "lastVisit": last_visit.isoformat() if last_visit else "",
            "daysSinceVisit": days_since_visit,
            "branch": branch,
            "advisor": advisor,
            "salesRepresentative": sales_representative,
            "technician": technician,
            "memberPassport": base.clean_text(t(14)),
            "ownerAuthorization": base.clean_text(t(13)),
            "spend3m": round(spend_3m, 0) if spend_3m is not None else None,
            "returnCount1y": base.as_int(t(17)),
            "maintenanceCount1y": base.as_int(t(17)),
            "returnCategory": base.clean_text(t(18)),
            "maintenanceFrequency": base.clean_text(t(19)),
            "beautySpend1y": base.as_int(t(21)),
            "beautyHighSpend": base.clean_text(t(22)),
            "fakeBeautySpend1y": base.as_int(t(23)),
            "windshieldCoatingCount1y": base.as_int(t(24)), "windshieldCoatingSpend1y": base.as_int(t(25)),
            "waxCount1y": base.as_int(t(26)), "waxSpend1y": base.as_int(t(27)),
            "bodyBeautyCount1y": base.as_int(t(28)), "bodyBeautySpend1y": base.as_int(t(29)),
            "interiorCleaningCount1y": base.as_int(t(30)), "interiorCleaningSpend1y": base.as_int(t(31)),
            "bodyCoatingCount1y": base.as_int(t(32)), "bodyCoatingSpend1y": base.as_int(t(33)),
            "warrantyStatus": warranty_status,
            "warrantyExpiry": warranty_expiry.isoformat() if warranty_expiry else "",
            "warrantyDaysLeft": days_left,
            "warrantyKmLeft": km_left,
            "extendedWarrantyExpiry": base.iso(row[62]),
            "tireChanges1y": tires, "batteryChanges1y": battery,
            "airFilterChanges1y": base.as_int(t(36)), "cabinFilterChanges1y": base.as_int(t(37)),
            "wiperChanges1y": base.as_int(t(38)), "brakePadChanges1y": base.as_int(t(39)),
            "beltChanges1y": base.as_int(t(40)), "sparkPlugChanges1y": base.as_int(t(41)),
            "tireLastChanged": base.iso(t(42)), "batteryLastChanged": base.iso(t(43)),
            "airFilterLastChanged": base.iso(t(44)), "cabinFilterLastChanged": base.iso(t(45)),
            "wiperLastChanged": base.iso(t(46)), "brakePadLastChanged": base.iso(t(47)),
            "beltLastChanged": base.iso(t(48)), "sparkPlugLastChanged": base.iso(t(49)),
            "contactAuthorized": base.clean_text(row[26]),
            "doNotContact": bool(base.clean_text(row[0])),
            "nextBestAction": base.recommended_action(warranty_status, days_since_visit, mileage, tires, battery),
        })
        for counter, value in ((models, model), (counties, county), (branches, branch), (advisors, advisor), (technicians, technician), (sales_representatives, sales_representative)):
            if value:
                counter[value] += 1
        warranty_counts[warranty_status] += 1

    workbook.close()
    if leaks:
        raise SystemExit(f"Refusing to write: non-code values found in personal-data columns {dict(leaks)}")

    with (OUTPUT_DIR / "vehicles.json").open("w", encoding="utf-8") as handle:
        json.dump(records, handle, ensure_ascii=False, separators=(",", ":"))
    metadata = {
        "source": SOURCE_LABEL,
        "asOf": AS_OF.isoformat(),
        "deidentified": True,
        "recordCount": len(records),
        "models": [name for name, _ in models.most_common()],
        "counties": [name for name, _ in counties.most_common()],
        "branches": [name for name, _ in branches.most_common()],
        "advisors": [name for name, _ in advisors.most_common()],
        "technicians": [name for name, _ in technicians.most_common()],
        "salesRepresentatives": [name for name, _ in sales_representatives.most_common()],
        "warrantyCounts": dict(warranty_counts),
    }
    with (OUTPUT_DIR / "metadata.json").open("w", encoding="utf-8") as handle:
        json.dump(metadata, handle, ensure_ascii=False, indent=2)
    print(json.dumps({"records": len(records), "duplicateCodes": duplicates, "withoutSummaryRow": unmatched, "bytes": (OUTPUT_DIR / "vehicles.json").stat().st_size, "warranty": dict(warranty_counts)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
